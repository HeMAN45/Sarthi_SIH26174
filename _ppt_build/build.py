# -*- coding: utf-8 -*-
"""Fill the official SIH 2026 template for ORBITAL-HAR (SIH26174, Team Hashira)."""
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn
import copy

SRC = r"C:\dse_3\SIH2026-IDEA-Presentation-Format.pptx"
OUT = r"C:\SISH_F\Clippy\ORBITAL-HAR_SIH26174_Hashira.pptx"

NAVY = RGBColor(0x1F, 0x49, 0x7D)
BLUE = RGBColor(0x00, 0x70, 0xC0)
LBLUE = RGBColor(0xE8, 0xF0, 0xFA)
MBLUE = RGBColor(0xD3, 0xE3, 0xF6)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
DARK = RGBColor(0x26, 0x2A, 0x33)
GREY = RGBColor(0x80, 0x80, 0x80)
YEL = RGBColor(0xFF, 0xF2, 0xCC)
YELB = RGBColor(0xD8, 0xA0, 0x00)
RED = RGBColor(0xC0, 0x39, 0x2B)
GREEN = RGBColor(0x2E, 0x7D, 0x32)
LGREY = RGBColor(0xF0, 0xF2, 0xF5)

prs = Presentation(SRC)
slides = prs.slides


def shape_by_name(slide, name):
    for sh in slide.shapes:
        if sh.name == name:
            return sh
    return None


def no_bullet(p):
    """Suppress any inherited template bullet on this paragraph."""
    pPr = p._p.get_or_add_pPr()
    for tag in ("a:buChar", "a:buAutoNum", "a:buNone"):
        for el in pPr.findall(qn(tag)):
            pPr.remove(el)
    pPr.append(pPr.makeelement(qn("a:buNone"), {}))


def set_title_native(slide, name, text):
    """Replace a title's text but keep the template's native font/size/colour."""
    sh = shape_by_name(slide, name)
    if sh is None:
        return
    tf = sh.text_frame
    tf.clear()
    p = tf.paragraphs[0]
    p.add_run().text = text


def set_para(p, text, size, color=DARK, bold=False, bullet=False,
             space_after=6, font="Calibri", align=PP_ALIGN.LEFT, italic=False):
    p.alignment = align
    p.space_after = Pt(space_after)
    p.space_before = Pt(0)
    no_bullet(p)
    run = p.add_run()
    run.text = ("\u2022  " + text) if bullet else text
    f = run.font
    f.size = Pt(size)
    f.bold = bold
    f.italic = italic
    f.name = font
    f.color.rgb = color
    return p


def fill_textbox(shape, items, size=14, space_after=7):
    """items: list of dicts {text,size?,color?,bold?,bullet?,italic?}"""
    tf = shape.text_frame
    tf.word_wrap = True
    tf.clear()
    first = True
    for it in items:
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        set_para(p, it["text"], it.get("size", size), it.get("color", DARK),
                 it.get("bold", False), it.get("bullet", False),
                 it.get("space_after", space_after), it.get("font", "Calibri"),
                 it.get("align", PP_ALIGN.LEFT), it.get("italic", False))


def set_oval_team(slide, oval_name, team):
    sh = shape_by_name(slide, oval_name)
    if sh is None:
        return
    tf = sh.text_frame
    tf.clear()
    p = tf.paragraphs[0]
    no_bullet(p)
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run(); r.text = team
    r.font.size = Pt(12); r.font.bold = True; r.font.name = "Calibri"
    r.font.color.rgb = NAVY


# ---------------- SLIDE 1 : TITLE ----------------
s1 = slides[0]
sub = shape_by_name(s1, "Subtitle 3")
if sub:
    # sit the project name clearly below the "SMART INDIA HACKATHON 2026" title;
    # inherit the template's native (serif) title font — set size/colour only.
    sub.top = Inches(1.55); sub.left = Inches(0.36); sub.width = Inches(9.6); sub.height = Inches(1.0)
    tf = sub.text_frame; tf.word_wrap = True; tf.clear()
    p = tf.paragraphs[0]; no_bullet(p); p.alignment = PP_ALIGN.LEFT; p.space_after = Pt(2)
    r = p.add_run(); r.text = "ORBITAL-HAR"
    r.font.size = Pt(32); r.font.bold = True; r.font.color.rgb = NAVY
    p2 = tf.add_paragraph(); no_bullet(p2); p2.alignment = PP_ALIGN.LEFT
    r2 = p2.add_run(); r2.text = "Offline on-board procedure supervision for BAS experiments"
    r2.font.size = Pt(15); r2.font.italic = True; r2.font.color.rgb = BLUE

tb = shape_by_name(s1, "TextBox 9")
if tb:
    tb.top = Inches(2.7); tb.left = Inches(0.36); tb.width = Inches(6.6); tb.height = Inches(4.4)
if tb:
    rows = [
        ("Problem Statement ID", "SIH26174"),
        ("Problem Statement Title", "AI Human Activity Recognition for On-board BAS Experiments"),
        ("Theme", "Space Technology"),
        ("PS Category", "Software"),
        ("Team ID", "(as registered on SIH portal)"),
        ("Team Name", "Hashira"),
    ]
    tf = tb.text_frame; tf.word_wrap = True; tf.clear()
    first = True
    for label, val in rows:
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.space_after = Pt(9); p.alignment = PP_ALIGN.LEFT
        r = p.add_run(); r.text = label + ":  "
        r.font.size = Pt(15); r.font.bold = True; r.font.name = "Calibri"; r.font.color.rgb = NAVY
        r2 = p.add_run(); r2.text = val
        r2.font.size = Pt(15); r2.font.name = "Calibri"; r2.font.color.rgb = DARK

# ---------------- SLIDE 2 : IDEA TITLE / PROPOSED SOLUTION ----------------
s2 = slides[1]
set_title_native(s2, "Title 1", "ORBITAL-HAR")
set_oval_team(s2, "Oval 9", "Hashira")
box2 = shape_by_name(s2, "TextBox 8")
# enlarge content region for the richer text
box2.top = Inches(1.55); box2.left = Inches(0.5); box2.width = Inches(12.3); box2.height = Inches(5.2)
fill_textbox(box2, [
    {"text": "The idea in one line", "size": 15, "bold": True, "color": BLUE, "space_after": 3},
    {"text": "A fixed payload camera watches an astronaut run an experiment. The system recognises each action, prompts the next step, voice-alerts on a skipped or out-of-order step, and emits a few KB of verified telemetry instead of a video downlink \u2014 fully offline on the edge.",
     "size": 13.5, "space_after": 10},
    {"text": "How it addresses the problem", "size": 15, "bold": True, "color": BLUE, "space_after": 3},
    {"text": "Replaces high-bandwidth video downlink with tamper-evident, hash-chained step telemetry \u2014 ground control gets provable status without real-time comms.", "size": 13, "bullet": True, "space_after": 5},
    {"text": "Reasoning-first, never guesses: perception publishes primitives (objects, hands, pose) to an event bus; a deterministic state machine decides and, below confidence, reports UNVERIFIED rather than a confident wrong verdict.", "size": 13, "bullet": True, "space_after": 5},
    {"text": "Orientation-agnostic: the input frame is canonicalised to the payload rack (ArUco), so a floating or inverted astronaut is still recognised \u2014 solving the \u201cno fixed up/down\u201d problem.", "size": 13, "bullet": True, "space_after": 10},
    {"text": "Innovation & uniqueness", "size": 15, "bold": True, "color": BLUE, "space_after": 3},
    {"text": "Procedure-as-config: every experiment is a YAML file, not code \u2014 a new experiment is a new file, with zero changes to the engine.", "size": 13, "bullet": True, "space_after": 5},
    {"text": "Event-bus seam decouples AI from reasoning: enables record-once/replay-1000\u00d7 tuning and \u201cdemo insurance\u201d if a camera fails live.", "size": 13, "bullet": True, "space_after": 5},
    {"text": "Explainable + auditable: deterministic verdicts and an append-only hash-chained log make every decision defensible.", "size": 13, "bullet": True, "space_after": 0},
])

# ---------------- SLIDE 3 : TECHNICAL APPROACH ----------------
s3 = slides[2]
set_title_native(s3, "Title 1", "TECHNICAL APPROACH")
set_oval_team(s3, "Oval 10", "Hashira")
box3 = shape_by_name(s3, "TextBox 8")
box3.top = Inches(1.2); box3.left = Inches(0.5); box3.width = Inches(12.3); box3.height = Inches(1.25)
fill_textbox(box3, [
    {"text": "Technologies", "size": 13.5, "bold": True, "color": BLUE, "space_after": 2},
    {"text": "Python 3.11  \u00b7  YOLO11 detection (object states as classes)  \u00b7  MediaPipe Hands  \u00b7  YOLO11-pose  \u00b7  OpenCV ArUco (rack frame)  \u00b7  Pydantic + YAML procedures  \u00b7  Piper offline TTS  \u00b7  FastAPI + WebSocket  \u00b7  React + Vite (bundled, no CDN)  \u00b7  FFmpeg + MediaMTX (RTSP + local record)  \u00b7  Blender synthetic data  \u00b7  Jetson Orin Nano + TensorRT (edge)",
     "size": 11.5, "space_after": 0},
])

# ----- architecture diagram on slide 3 -----
def box(slide, x, y, w, h, text, fill, line, txt=DARK, size=10.5, bold=True,
        shape=MSO_SHAPE.ROUNDED_RECTANGLE, line_w=1.25):
    sp = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    sp.fill.solid(); sp.fill.fore_color.rgb = fill
    if line is None:
        sp.line.fill.background()
    else:
        sp.line.color.rgb = line; sp.line.width = Pt(line_w)
    sp.shadow.inherit = False
    tf = sp.text_frame; tf.word_wrap = True
    tf.margin_left = Inches(0.04); tf.margin_right = Inches(0.04)
    tf.margin_top = Inches(0.02); tf.margin_bottom = Inches(0.02)
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    lines = text.split("\n")
    for i, ln in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = PP_ALIGN.CENTER; p.space_after = Pt(0); p.space_before = Pt(0)
        r = p.add_run(); r.text = ln
        r.font.size = Pt(size if i == 0 else size - 1.5)
        r.font.bold = bold if i == 0 else False
        r.font.name = "Calibri"; r.font.color.rgb = txt
    return sp

def arrow(slide, x, y, w, h, color=BLUE, shape=MSO_SHAPE.RIGHT_ARROW):
    sp = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    sp.fill.solid(); sp.fill.fore_color.rgb = color
    sp.line.fill.background(); sp.shadow.inherit = False
    return sp

# main pipeline row
ROW_Y = 3.05
BH = 0.95
BW = 2.05
GAP = 0.30
xs = 0.45
cols = []
for i in range(5):
    cols.append(xs + i * (BW + GAP))

# 1 inputs (stacked labels in one box)
box(s3, cols[0], ROW_Y, BW, BH,
    "Payload Camera\n+ Rack ArUco Markers", LBLUE, BLUE, DARK, 11)
# 2 canonicalizer
box(s3, cols[1], ROW_Y, BW, BH,
    "Rack-frame\nCanonicalizer", MBLUE, BLUE, DARK, 11)
# 3 perception
box(s3, cols[2], ROW_Y, BW, BH,
    "Perception\nDetect \u00b7 Hands \u00b7 Pose", MBLUE, BLUE, DARK, 11)
# 4 event bus (highlight seam)
box(s3, cols[3], ROW_Y, BW, BH,
    "Event Bus\n(JSONL stream)", BLUE, BLUE, WHITE, 11.5)
# 5 engine (brain)
box(s3, cols[4], ROW_Y, BW, BH,
    "Procedure Engine\ndeterministic state machine", NAVY, NAVY, WHITE, 11)

# arrows between top-row boxes
for i in range(4):
    ax = cols[i] + BW + 0.01
    arrow(s3, ax, ROW_Y + BH/2 - 0.11, GAP - 0.02, 0.22, BLUE)

# procedure YAML feeding the engine from below (offset left so it clears the outputs drop)
yaml_y = 4.55
box(s3, 8.30, yaml_y, 2.60, 0.62,
    "Procedure YAML\n(experiment as config)", YEL, YELB, DARK, 10.5)
arrow(s3, 10.15, ROW_Y + BH + 0.01, 0.22, yaml_y - (ROW_Y + BH) - 0.01,
      YELB, MSO_SHAPE.UP_ARROW)

# outputs row (4 boxes)
OUT_Y = 5.80
oW = 2.95; oGAP = 0.20; oH = 0.80; oxs = 0.55
step = oW + oGAP
centers = [oxs + i * step + oW / 2 for i in range(4)]

# distribution manifold: engine -> bar -> 4 outputs
bar_y = 5.42
bar_x0 = centers[0]; bar_x1 = centers[3]
barsp = s3.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(bar_x0), Inches(bar_y),
                            Inches(bar_x1 - bar_x0), Inches(0.05))
barsp.fill.solid(); barsp.fill.fore_color.rgb = NAVY
barsp.line.fill.background(); barsp.shadow.inherit = False
# engine feeds the manifold (drop from the engine's right edge, clear of the YAML link)
eng_cx = 11.45
arrow(s3, eng_cx - 0.11, ROW_Y + BH + 0.02, 0.22, bar_y - (ROW_Y + BH) - 0.02,
      NAVY, MSO_SHAPE.DOWN_ARROW)
# manifold drops into each output
for cx in centers:
    arrow(s3, cx - 0.09, bar_y + 0.05, 0.18, OUT_Y - bar_y - 0.06, NAVY, MSO_SHAPE.DOWN_ARROW)

outs = [
    ("Voice Alerts\nPiper TTS (offline)", RED),
    ("Hash-chained\nTelemetry (KB, auditable)", GREEN),
    ("RTSP Video Out\n+ Local Recording", BLUE),
    ("Dashboard\nCrew HUD / Ground Ops", NAVY),
]
for i, (txt, c) in enumerate(outs):
    ox = oxs + i * step
    box(s3, ox, OUT_Y, oW, oH, txt, LGREY, c, DARK, 10.5, line_w=1.5)

# caption (bottom, above footer)
cap = s3.shapes.add_textbox(Inches(0.55), Inches(6.66), Inches(11.8), Inches(0.28))
cp = cap.text_frame.paragraphs[0]
rr = cp.add_run(); rr.text = "Fire-and-forget outputs \u2014 losing any output never stops supervision."
rr.font.size = Pt(10); rr.font.italic = True; rr.font.color.rgb = GREY; rr.font.name = "Calibri"

# ---------------- SLIDE 4 : FEASIBILITY AND VIABILITY ----------------
s4 = slides[3]
set_title_native(s4, "Title 1", "FEASIBILITY AND VIABILITY")
set_oval_team(s4, "Oval 11", "Hashira")
box4 = shape_by_name(s4, "TextBox 8")
box4.top = Inches(1.55); box4.left = Inches(0.5); box4.width = Inches(12.3); box4.height = Inches(5.2)
fill_textbox(box4, [
    {"text": "Feasibility \u2014 already de-risked", "size": 15, "bold": True, "color": BLUE, "space_after": 3},
    {"text": "The reasoning spine is built and tested end-to-end with no camera and no models: procedure schema, event bus with replay, 7 predicate evaluators, step state machine and CLI \u2014 73 tests pass, plus a golden replay corpus.", "size": 13, "bullet": True, "space_after": 5},
    {"text": "A live webcam demo already runs today using a pretrained YOLO stand-in feeding the real engine (on-screen HUD + offline voice).", "size": 13, "bullet": True, "space_after": 5},
    {"text": "Small-data regime: fixed camera, controlled scene, under 12 classes \u2014 ~2,000 synthetic + ~800 real frames is enough to fine-tune YOLO11 from COCO weights.", "size": 13, "bullet": True, "space_after": 10},
    {"text": "Challenges & risks", "size": 15, "bold": True, "color": BLUE, "space_after": 3},
    {"text": "Dataset generation is the critical path (props \u2192 footage \u2192 labels \u2192 model). Standard pose models fail on inverted bodies. Edge compute/power budget and real-time alert latency are constrained.", "size": 13, "bullet": True, "space_after": 10},
    {"text": "Strategies to overcome them", "size": 15, "bold": True, "color": BLUE, "space_after": 3},
    {"text": "Blender auto-labelled synthetic data randomised across the full orientation sphere; canonicalise the input frame (not the output pose); pre-synthesize voice prompts to WAV off the alert path; TensorRT export for the Jetson; event-bus replay as live demo insurance.", "size": 13, "bullet": True, "space_after": 0},
])

# ---------------- SLIDE 5 : IMPACT AND BENEFITS ----------------
s5 = slides[4]
set_title_native(s5, "Title 1", "IMPACT AND BENEFITS")
set_oval_team(s5, "Oval 11", "Hashira")
box5 = shape_by_name(s5, "TextBox 8")
box5.top = Inches(1.55); box5.left = Inches(0.5); box5.width = Inches(12.3); box5.height = Inches(5.2)
fill_textbox(box5, [
    {"text": "Mission impact", "size": 15, "bold": True, "color": BLUE, "space_after": 3},
    {"text": "Enables autonomous, protocol-perfect science execution beyond real-time ground support (BAS, lunar) \u2014 the core need in the problem statement.", "size": 13, "bullet": True, "space_after": 8},
    {"text": "Bandwidth saved", "size": 15, "bold": True, "color": BLUE, "space_after": 3},
    {"text": "KB-scale verified telemetry instead of GB-scale video downlink \u2014 orders-of-magnitude saving on a data-restricted space link.", "size": 13, "bullet": True, "space_after": 8},
    {"text": "Safety & trust", "size": 15, "bold": True, "color": BLUE, "space_after": 3},
    {"text": "Deterministic, explainable verdicts and a tamper-evident hash-chained record of every step reduce crew error and support post-mission audit.", "size": 13, "bullet": True, "space_after": 8},
    {"text": "Reusability & spin-off", "size": 15, "bold": True, "color": BLUE, "space_after": 3},
    {"text": "Procedure-as-config makes it a platform, not a one-experiment demo \u2014 any procedural task (surgical, industrial, hazardous-site, training) reuses the same engine.", "size": 13, "bullet": True, "space_after": 8},
    {"text": "Economic & social", "size": 15, "bold": True, "color": BLUE, "space_after": 3},
    {"text": "Lowers ground-ops load and crew training burden; indigenous, offline edge-AI capability aligned with India's space programme.", "size": 13, "bullet": True, "space_after": 0},
])

# ---------------- SLIDE 6 : RESEARCH AND REFERENCES ----------------
s6 = slides[5]
set_title_native(s6, "Title 1", "RESEARCH AND REFERENCES")
set_oval_team(s6, "Oval 8", "Hashira")
box6 = shape_by_name(s6, "TextBox 8")
box6.top = Inches(1.7); box6.left = Inches(0.5); box6.width = Inches(12.3); box6.height = Inches(5.0)
fill_textbox(box6, [
    {"text": "SIH 2026 Problem Statement SIH26174 \u2014 AI HAR for On-board BAS Experiments, ISRO / Department of Space  (sih.gov.in/sih2026PS)", "size": 13, "bullet": True, "space_after": 7},
    {"text": "Ultralytics YOLO11 / YOLO-World \u2014 object detection & open-vocabulary pre-labelling  (docs.ultralytics.com)", "size": 13, "bullet": True, "space_after": 7},
    {"text": "Google MediaPipe Hands \u2014 real-time hand landmarks & hand-object interaction", "size": 13, "bullet": True, "space_after": 7},
    {"text": "OpenCV ArUco \u2014 6-DoF fiducial marker pose for rack-frame canonicalisation", "size": 13, "bullet": True, "space_after": 7},
    {"text": "Piper TTS (Rhasspy) \u2014 fully offline neural text-to-speech for voice alerts", "size": 13, "bullet": True, "space_after": 7},
    {"text": "Orientation-agnostic 3D Human Mesh Recovery (HMR) literature \u2014 rack-relative body tracking", "size": 13, "bullet": True, "space_after": 7},
    {"text": "Blender bpy \u2014 scripted synthetic-data generation with automatic labels", "size": 13, "bullet": True, "space_after": 0},
])

# ---------------- delete slide 7 (instructions) ----------------
# Remove the sldId entry AND drop the relationship so the physical part,
# its rels and its content-type override are not written to the package.
xml_slides = prs.slides._sldIdLst
sld_list = list(xml_slides)
rid = sld_list[6].get(qn("r:id"))
xml_slides.remove(sld_list[6])
prs.part.drop_rel(rid)

prs.save(OUT)
print("SAVED:", OUT)
print("slides:", len(prs.slides._sldIdLst))
